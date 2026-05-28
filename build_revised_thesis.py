from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE_DOCX = ROOT / "基于树莓派的婴幼儿智能陪护小车系统设计.docx"
OUTPUT_DOCX = ROOT / "基于树莓派的婴幼儿智能陪护小车系统设计_严格格式版_封面按规范图.docx"
FIG_DIR = ROOT / "docs" / "thesis_figures"
SCHEMATIC = ROOT / "文献" / "SCH_外围电路原理图_1-1_2026-05-18.png"


def set_run_font(run, font="宋体", size=12, bold=False, italic=False):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)


def clear_document(doc: Document) -> None:
    body = doc._body._element
    for child in list(body):
        if child.tag.endswith("sectPr"):
            continue
        body.remove(child)


def configure_doc(doc: Document) -> None:
    for section in doc.sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(3.175)
        section.right_margin = Cm(3.175)
    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.first_line_indent = Cm(0.74)
    for name in ("Heading 1", "Heading 2", "Heading 3", "标题 1", "标题 2", "标题 3"):
        if name in doc.styles:
            style = doc.styles[name]
            style.font.color.rgb = RGBColor(0, 0, 0)
            style.font.name = "黑体"
            style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    if "Caption" in doc.styles:
        cap = doc.styles["Caption"]
        cap.font.name = "黑体"
        cap._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        cap.font.size = Pt(10.5)
        cap.font.bold = True
        cap.font.color.rgb = RGBColor(0, 0, 0)


def para(doc, text="", indent=True, style=None, font="宋体", size=12, bold=False):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.first_line_indent = Cm(0.74) if indent else Cm(0)
    if text:
        run = p.add_run(text)
        set_run_font(run, font=font, size=size, bold=bold)
    return p


def centered(doc, text, font="黑体", size=16, bold=True):
    p = para(doc, "", indent=False)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, font=font, size=size, bold=bold)
    return p


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.first_line_indent = Cm(0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
    size = {1: 16, 2: 14, 3: 12}.get(level, 12)
    for run in p.runs:
        set_run_font(run, font="黑体", size=size, bold=True)
    return p


def caption(doc, text):
    p = para(doc, text, indent=False, style="Caption" if "Caption" in doc.styles else None, font="黑体", size=10.5, bold=True)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p


def _set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = tcPr.first_child_found_in("w:tcBorders")
    if tcBorders is None:
        tcBorders = OxmlElement("w:tcBorders")
        tcPr.append(tcBorders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge in kwargs:
            tag = "w:{}".format(edge)
            element = tcBorders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                tcBorders.append(element)
            for key, value in kwargs[edge].items():
                element.set(qn("w:{}".format(key)), str(value))


def set_three_line_table(table):
    nil = {"val": "nil", "sz": "0", "space": "0", "color": "FFFFFF"}
    thin = {"val": "single", "sz": "8", "space": "0", "color": "000000"}
    thick = {"val": "single", "sz": "12", "space": "0", "color": "000000"}
    for row in table.rows:
        for cell in row.cells:
            _set_cell_border(cell, top=nil, left=nil, bottom=nil, right=nil, insideH=nil, insideV=nil)
    for cell in table.rows[0].cells:
        _set_cell_border(cell, top=thick, bottom=thin, left=nil, right=nil)
    for cell in table.rows[-1].cells:
        _set_cell_border(cell, bottom=thick, left=nil, right=nil)


def add_table(doc, headers, rows, widths=None, title=None):
    if title:
        caption(doc, title)
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(str(h))
        set_run_font(run, font="黑体", size=12, bold=True)
        if widths:
            cell.width = Cm(widths[i])
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(str(value))
            set_run_font(run, size=12)
            if widths:
                cells[i].width = Cm(widths[i])
    set_three_line_table(table)
    para(doc, "", indent=False)
    return table


def cover_line(doc, label, value, left_cm=2.26, size=16):
    p = para(doc, "", indent=False)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.left_indent = Cm(left_cm)
    p.paragraph_format.first_line_indent = Cm(0)
    run = p.add_run(f"{label}{value}")
    set_run_font(run, size=size, bold=True)
    return p


def underlined_cover_line(doc, label, value, left_cm=3.55, label_width=5, line_chars=22, value_font="Times New Roman"):
    p = para(doc, "", indent=False)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.left_indent = Cm(left_cm)
    p.paragraph_format.first_line_indent = Cm(0)
    r_label = p.add_run(label)
    set_run_font(r_label, font="宋体", size=16, bold=True)
    r_value = p.add_run(value.center(line_chars))
    set_run_font(r_value, font=value_font, size=16, bold=True)
    r_value.font.underline = True
    return p


def add_figure(doc, path: Path | None, title: str, width=5.8, placeholder=None):
    if path and path.exists():
        p = para(doc, "", indent=False)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(str(path), width=Inches(width))
    elif placeholder:
        p = para(doc, f"[{placeholder}]", indent=False, size=10.5)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption(doc, title)


def add_page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def _add_field(run, field):
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(fld_end)


def _set_page_numbering(section, start=1, fmt="decimal"):
    sectPr = section._sectPr
    pgNumType = sectPr.find(qn("w:pgNumType"))
    if pgNumType is None:
        pgNumType = OxmlElement("w:pgNumType")
        sectPr.append(pgNumType)
    pgNumType.set(qn("w:start"), str(start))
    pgNumType.set(qn("w:fmt"), fmt)


def configure_section(section, header=True, page_number=True, start=None, number_format="decimal"):
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.18)
    section.right_margin = Cm(3.18)
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    for p in section.header.paragraphs:
        p.clear()
    for p in section.footer.paragraphs:
        p.clear()
    if header:
        hp = section.header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = hp.add_run("西安工程大学本科毕业设计（论文）")
        set_run_font(r, size=10.5)
    if page_number:
        fp = section.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = fp.add_run()
        set_run_font(r, font="Times New Roman", size=10.5)
        _add_field(r, "PAGE")
    if start is not None:
        _set_page_numbering(section, start=start, fmt=number_format)


def new_section(doc, header=True, page_number=True, start=None, number_format="decimal"):
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(section, header=header, page_number=page_number, start=start, number_format=number_format)
    return section


def make_diagrams():
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    try:
        from PIL import Image, ImageDraw, ImageFont
    except Exception:
        return

    font_candidates = [
        Path(r"C:\Windows\Fonts\msyh.ttc"),
        Path(r"C:\Windows\Fonts\simsun.ttc"),
        Path(r"C:\Windows\Fonts\simhei.ttf"),
    ]
    font_path = next((p for p in font_candidates if p.exists()), None)
    font = ImageFont.truetype(str(font_path), 26) if font_path else ImageFont.load_default()
    small = ImageFont.truetype(str(font_path), 20) if font_path else ImageFont.load_default()

    def save_boxes(name, title, boxes, arrows):
        img = Image.new("RGB", (1400, 760), "white")
        d = ImageDraw.Draw(img)
        d.text((700, 30), title, anchor="mm", fill=(20, 20, 20), font=font)
        colors = [(226, 241, 255), (233, 246, 232), (255, 242, 222), (244, 235, 255)]
        coords = {}
        for idx, (key, label, x, y, w, h) in enumerate(boxes):
            fill = colors[idx % len(colors)]
            d.rounded_rectangle((x, y, x + w, y + h), radius=18, fill=fill, outline=(80, 110, 140), width=3)
            lines = label.split("\n")
            for j, line in enumerate(lines):
                d.text((x + w / 2, y + h / 2 - 14 * (len(lines) - 1) + j * 28), line, anchor="mm", fill=(0, 0, 0), font=small)
            coords[key] = (x, y, w, h)
        for a, b, label in arrows:
            ax, ay, aw, ah = coords[a]
            bx, by, bw, bh = coords[b]
            start = (ax + aw, ay + ah / 2)
            end = (bx, by + bh / 2)
            if bx < ax:
                start = (ax, ay + ah / 2)
                end = (bx + bw, by + bh / 2)
            d.line((start, end), fill=(45, 85, 120), width=4)
            d.polygon([(end[0], end[1]), (end[0] - 14, end[1] - 8), (end[0] - 14, end[1] + 8)], fill=(45, 85, 120))
            if label:
                d.text(((start[0] + end[0]) / 2, (start[1] + end[1]) / 2 - 24), label, anchor="mm", fill=(50, 50, 50), font=small)
        img.save(FIG_DIR / name)

    save_boxes(
        "fig2-1-architecture.png",
        "系统总体架构图",
        [
            ("car", "小车端\n摄像头/传感器/执行器", 90, 260, 300, 150),
            ("pc", "PC计算端\nYOLO/YAMNet/ASR/桥接", 550, 260, 320, 150),
            ("app", "Android端\n视频/环境/控制/报警", 1030, 260, 300, 150),
        ],
        [("car", "pc", "视频0x01/环境0x03"), ("pc", "app", "视频/状态/报警"), ("app", "pc", "控制0x02"), ("pc", "car", "控制0x02")],
    )
    save_boxes(
        "fig2-2-protocol-flow.png",
        "通信协议数据流图",
        [
            ("video", "0x01 视频帧\nJPEG 640x480", 70, 160, 270, 120),
            ("cmd", "0x02 控制指令\nJSON动作/云台/语音", 70, 340, 270, 120),
            ("env", "0x03 环境数据\nJSON传感器/报警", 70, 520, 270, 120),
            ("pc", "PC网关\n认证/转发/桥接", 560, 320, 300, 150),
            ("app", "Android App\n显示/控制/通知", 1050, 320, 280, 150),
        ],
        [("video", "pc", ""), ("cmd", "pc", ""), ("env", "pc", ""), ("pc", "app", "WebSocket/WebRTC"), ("app", "pc", "指令回传")],
    )
    save_boxes(
        "fig4-6-voice-flow.png",
        "语音交互流程图",
        [
            ("mic", "树莓派麦克风\nPCM 16kHz/16bit", 60, 300, 250, 130),
            ("asr", "PC AsrServer\n百度实时ASR", 390, 300, 250, 130),
            ("intent", "意图解析\n8类指令映射", 720, 300, 250, 130),
            ("exec", "动作执行/TTS反馈\n电机/儿歌/播报", 1050, 300, 280, 130),
        ],
        [("mic", "asr", "WebSocket音频"), ("asr", "intent", "识别文本"), ("intent", "exec", "控制/回复")],
    )


def add_cover_and_front_matter(doc):
    configure_section(doc.sections[0], header=False, page_number=False)
    for _ in range(1):
        para(doc, "", indent=False)
    p = para(doc, "", indent=False)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("西安工程大学")
    set_run_font(r, font="华文行楷", size=30, bold=False)
    for _ in range(2):
        para(doc, "", indent=False)
    p = para(doc, "", indent=False)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("毕业设计（论文）")
    set_run_font(r, font="宋体", size=16, bold=True)
    for _ in range(9):
        para(doc, "", indent=False)
    underlined_cover_line(doc, "题    目：", "基于树莓派的婴幼儿智能陪护小车系统设计", left_cm=3.45, line_chars=24, value_font="宋体")
    underlined_cover_line(doc, "学    院：", "电子信息学院", left_cm=3.45, line_chars=24, value_font="宋体")
    underlined_cover_line(doc, "专业班级：", "机器人工程2022级1班", left_cm=3.45, line_chars=23, value_font="Times New Roman")
    p = para(doc, "", indent=False)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.left_indent = Cm(3.45)
    p.paragraph_format.first_line_indent = Cm(0)
    r = p.add_run("指导教师：")
    set_run_font(r, font="宋体", size=16, bold=True)
    r = p.add_run("  张  蕾  ")
    set_run_font(r, font="宋体", size=16, bold=True)
    r.font.underline = True
    r = p.add_run("  职称：")
    set_run_font(r, font="宋体", size=16, bold=True)
    r = p.add_run("  教授  ")
    set_run_font(r, font="宋体", size=16, bold=True)
    r.font.underline = True
    underlined_cover_line(doc, "学生姓名：", "赵国羽", left_cm=3.45, line_chars=23, value_font="宋体")
    underlined_cover_line(doc, "学    号：", "42203080116", left_cm=3.45, line_chars=24, value_font="Times New Roman")
    new_section(doc, header=True, page_number=True, start=1, number_format="upperRoman")

    centered(doc, "摘  要", size=16)
    para(doc, "家庭婴幼儿看护中，看护者需要同时关注婴儿的活动状态、睡眠环境、哭闹信号以及潜在安全风险。当婴儿移动到固定摄像头视野之外时，看护者便失去关键视觉信息，这是传统固定式监护设备的重要局限。针对上述问题，本文设计并实现了一套基于树莓派的婴幼儿智能陪护小车系统，以移动小车为执行平台，融合视觉识别、哭声检测、环境监测、语音交互、远程视频查看和基础运动控制等功能。")
    para(doc, "系统采用小车端、PC计算端和Android移动端协同的三层结构。小车端完成传感器采集、运动执行、OLED显示、本地报警和安全保护；PC端完成目标识别、哭声判断、语音识别、控制决策和通信转发；Android端提供视频查看、环境状态显示、报警接收和远程控制入口。系统通过0x01视频帧、0x02控制指令和0x03环境数据三类消息统一三端通信，并支持局域网WebSocket与公网WebRTC两种远程访问方式。")
    para(doc, "系统关键技术包括：采用YOLO26s轻量化目标检测模型结合NONE/CANDIDATE/LOCKED三态BabyFilter时序滤波器实现婴幼儿目标稳定锁定；基于PID双环伺服控制实现云台视角调整与车身转向协同；基于YAMNet音频分类模型、RMS能量门控和双阈值滞后状态机实现实时哭声检测；通过百度实时语音识别、意图解析和TTS动态回声抑制构建语音交互闭环；基于距离约束、悬崖检测、通信看门狗、麦克风断连保护和控制权冲突管理构建多层级安全保护体系。实物调试结果表明，系统能够完成目标跟随、哭声提醒、环境异常提示、语音控制、移动端远程监控和安全保护等主要功能，验证了低成本移动平台用于婴幼儿辅助看护原型设计的可行性。")
    para(doc, "关键词：树莓派；婴幼儿陪护；PID双环伺服控制；YAMNet；语音交互；安全保护", bold=True)
    new_section(doc, header=True, page_number=True)

    centered(doc, "ABSTRACT", font="Times New Roman", size=16)
    para(doc, "Continuous observation, abnormal-event reminders and remote coordination are important requirements in home infant care. Conventional fixed monitoring devices usually provide only a single-view video or audio channel and are limited in active viewpoint adjustment, multi-source sensing and timely mobile feedback. To address these limitations, this thesis designs and implements a Raspberry Pi-based intelligent infant companion car that integrates visual recognition, cry detection, environmental monitoring, voice interaction, remote video viewing and basic motion control.", font="Times New Roman")
    para(doc, "The system is organized as a three-terminal architecture consisting of a car side, a PC computing side and an Android mobile side. The car side performs sensor acquisition, motion execution, OLED display, local alarm and safety protection. The PC side performs target recognition, cry judgement, speech recognition, control decision-making and communication forwarding. The Android side provides video viewing, environmental-status display, alarm reception and remote-control functions. Three message types, namely 0x01 video frames, 0x02 control commands and 0x03 environmental packets, are used to unify communication among terminals.", font="Times New Roman")
    para(doc, "Key techniques include YOLO26s-based target detection with a three-state BabyFilter, dual-loop PID servo control, YAMNet-based cry detection with a hysteresis state machine, Baidu real-time ASR with intent mapping and TTS echo suppression, and a multi-layer safety mechanism covering distance constraints, cliff detection, communication watchdogs, microphone health checks and control-priority arbitration. Physical debugging shows that the system can complete target following, cry reminders, environmental-abnormality prompts, voice control, mobile remote monitoring and safety protection, which verifies the feasibility of using a low-cost mobile platform for an infant-care assistance prototype.", font="Times New Roman")
    para(doc, "KEY WORDS: Raspberry Pi; infant companion; dual-loop PID; YAMNet; voice interaction; safety protection", bold=True, font="Times New Roman")
    new_section(doc, header=True, page_number=True)


def add_toc_field(paragraph):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = r'TOC \o "1-3" \h \z \u'
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "目录生成中，请在 Word 中更新域。"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(placeholder)
    run._r.append(fld_end)


def add_manual_toc(doc):
    centered(doc, "目  录", size=18)
    p = para(doc, "", indent=False)
    add_toc_field(p)
    new_section(doc, header=True, page_number=True, start=1, number_format="decimal")


def build():
    make_diagrams()
    doc = Document(SOURCE_DOCX)
    clear_document(doc)
    configure_doc(doc)
    add_cover_and_front_matter(doc)
    add_manual_toc(doc)

    heading(doc, "第1章 引言", 1)
    heading(doc, "1.1 研究背景与意义", 2)
    para(doc, "家庭婴幼儿看护中，看护者需要同时关注婴儿的活动状态、睡眠环境、哭闹信号以及潜在的安全风险。当婴儿移动到摄像头视野之外时，看护者便失去视觉信息，这是固定式监护设备最根本的局限。随着双职工家庭比例提高，家长在做家务、远程办公或短暂离开房间时，往往难以持续保持同一强度的看护，因此辅助看护设备需要具备连续观察、异常提醒、远程查看和安全干预能力。")
    para(doc, "与固定智能婴儿床不同，移动平台能够在室内主动跟随婴幼儿位置，保持更连续的观察视角。树莓派小车具备低成本、接口开放和二次开发便利等特点，适合将视觉识别、声音分析、环境监测、运动控制和移动端交互集成到同一原型系统中。本文围绕家庭室内辅助看护场景，设计基于树莓派的婴幼儿智能陪护小车系统，重点验证多源感知、主动视觉观察和三端远程协同的工程可行性。")

    heading(doc, "1.2 国内外研究现状", 2)
    para(doc, "现有婴幼儿监护与家庭陪护方案大致可分为固定摄像头监护、智能婴儿床、树莓派或STM32移动平台、无线传感监护和综合移动陪护系统等类型。固定摄像头婴儿监护器能够提供视频查看和夜视功能，但视角固定，无法在婴幼儿离开画面时主动调整观察方向；智能婴儿床系统通常具备环境监测、音乐安抚和App控制能力，但设备位置固定，主动观察和移动跟随能力不足。")
    para(doc, "基于树莓派的智能小车研究多集中在循迹、避障、远程控制和通用目标识别等任务，能够为移动平台设计提供参考，但缺少面向婴幼儿看护场景的专项安全策略和多模态协同设计。基于STM32的追踪监护方案具备低成本和实时控制优势，但受算力限制，复杂视觉识别、哭声检测和远程视频传输能力相对有限。基于ZigBee等无线传感网络的婴儿监护方案部署简单、功耗较低，但一般缺少视觉感知和移动执行能力。")
    para(doc, "相比上述方案，本文系统以移动小车为载体，将视觉识别、主动跟随、哭声检测、语音交互、环境监测、双模远程视频和多层级安全保护集成在同一原型系统中。已有研究在视觉感知、声音分析、环境监测、移动控制和远程交互之间相对分散，缺少将多源感知、主动观察、语音对话、安全移动和远程监控统一起来的移动陪护系统。本文据此设计基于树莓派的移动小车平台，面向家庭室内辅助看护场景开展系统设计与验证。")

    heading(doc, "1.3 本文主要工作", 2)
    for idx, item in enumerate([
        "设计三端协同总体架构。系统由树莓派小车端、PC计算端和Android App端组成，小车端负责感知执行，PC端负责智能计算与通信桥接，Android端负责状态呈现、远程控制和报警通知。",
        "完成硬件接口与资源规划。围绕树莓派40Pin、I2C总线、CSI摄像头和USB音频接口，完成超声波、PCF8591、MPU6050、红外循迹、OLED、蜂鸣器、云台舵机、麦克风和扬声器等模块集成。",
        "实现婴幼儿视觉跟随。PC端使用YOLO26s目标检测模型，结合BabyFilter三态时序滤波、PID双环伺服控制和距离跟随状态机，实现从图像识别到云台和底盘控制的闭环。",
        "实现哭声检测与语音交互。基于YAMNet和双阈值滞后状态机识别持续哭声，基于百度实时ASR和意图解析支持8类语音指令，并设计TTS回声抑制机制降低自激误触发。",
        "实现通信协议与远程监控。定义0x01视频、0x02指令和0x03环境数据三类消息，支持局域网WebSocket和公网WebRTC双模视频传输，并在Android端完成视频、环境数据、控制和报警展示。",
        "构建多层级安全保护机制。围绕距离约束、悬崖检测、通信看门狗、麦克风断连保护、运动指令前置检查和控制权冲突管理，提升移动平台在家庭环境中的运行安全性。",
    ], 1):
        para(doc, f"（{idx}）{item}")
    heading(doc, "1.4 论文结构安排", 2)
    para(doc, "本文共分为六章。第一章介绍研究背景、国内外研究现状和主要工作；第二章给出系统总体设计、创新点和通信协议；第三章压缩并归纳硬件设计内容，说明主控资源、传感器与执行器接口、外围电路和实物集成；第四章重点阐述软件设计，包括视觉跟随、哭声检测、语音交互、环境监测、远程通信、安全保护和Android端设计；第五章展示系统实物调试与功能验证；第六章总结全文并提出后续改进方向。")
    add_page_break(doc)

    heading(doc, "第2章 系统总体设计", 1)
    heading(doc, "2.1 设计目标与需求分析", 2)
    para(doc, "系统设计目标包括：实现对婴幼儿的视觉稳定跟踪与主动观察；实现实时哭声检测与语音对话交互；实现多传感器环境监测与分级报警；实现局域网/公网双模远程视频监控；构建多层级安全保护机制。")
    add_table(doc, ["编号", "功能需求", "说明", "优先级"], [
        ["F1", "视觉稳定跟踪", "识别婴幼儿目标并调整云台和车身位置", "高"],
        ["F2", "哭声检测报警", "持续哭声触发报警并同步到小车端和App端", "高"],
        ["F3", "语音交互", "支持运动控制、停止、儿歌播放等8类指令", "中"],
        ["F4", "环境监测", "采集距离、温度、光照、烟雾、音量、循迹和IMU数据", "高"],
        ["F5", "远程视频监控", "支持局域网WebSocket和公网WebRTC两种连接", "中"],
        ["F6", "安全保护", "障碍物、悬崖、断连和控制冲突状态下优先保护", "高"],
    ], title="表 2-1 系统功能需求")
    heading(doc, "2.2 系统总体架构", 2)
    para(doc, "系统采用小车端、PC计算端和Android移动端协同的三层架构。小车端承担传感器采集、运动执行、本地显示、音频播放和安全动作；PC端承担YOLO26s目标检测、BabyFilter时序滤波、PID控制决策、YAMNet哭声检测、百度ASR、WebSocket转发和WebRTC桥接；Android端提供实时视频、环境数据、报警通知和远程控制界面。")
    add_figure(doc, FIG_DIR / "fig2-1-architecture.png", "图 2-1 系统总体架构图", width=6.0)
    heading(doc, "2.3 系统设计创新点", 2)
    innovation_items = [
        ("移动平台主动视觉观察", "与固定摄像头或智能婴儿床不同，本系统以移动小车为载体，结合YOLO26s目标检测、三态BabyFilter时序滤波和PID双环伺服控制，使摄像头云台和车身能够同步跟随婴幼儿位置变化，实现主动视角调整。"),
        ("基于YAMNet的双阈值哭声检测", "系统利用AudioSet预训练的YAMNet音频分类模型，通过1.0s滑动窗口、RMS能量门控和双阈值滞后状态机实现实时哭声检测，较固定阈值或简单分类器方案更适合复杂家庭声音环境。"),
        ("语音对话交互闭环", "系统构建麦克风采集、百度实时ASR、意图解析、动作执行和TTS语音反馈的完整链路，支持前进、后退、左右转、停止、播放儿歌、下一首和停止播放等指令，并通过动态回声抑制避免机器人自发声造成误触发。"),
        ("局域网/公网双模视频传输", "局域网模式通过WebSocket直连传输JPEG帧，公网模式通过WebRTC、STUN/TURN和DataChannel完成视频与状态同步，兼顾家庭内部低延迟查看和外出远程监控需求。"),
        ("多层级安全保护体系", "系统从超声波距离限行、红外悬崖检测、通信看门狗、麦克风断连保护到控制权冲突管理，构建硬件、软件和通信三个层面协同的安全防护网络。"),
    ]
    for idx, (title, body) in enumerate(innovation_items, 1):
        para(doc, f"（{idx}）{title}。{body}")
    heading(doc, "2.4 通信协议设计", 2)
    para(doc, "系统定义统一二进制帧协议，每个WebSocket数据包由1字节类型标识和后续负载组成。0x01表示视频帧，负载为JPEG原始字节；0x02表示控制指令，负载为UTF-8 JSON；0x03表示环境数据，负载为UTF-8 JSON。协议数据流如图2-2所示。")
    add_figure(doc, FIG_DIR / "fig2-2-protocol-flow.png", "图 2-2 通信协议数据流图", width=6.0)
    add_table(doc, ["类型标识", "方向", "含义", "负载格式"], [
        ["0x01", "小车/PC → App", "视频帧", "JPEG原始字节"],
        ["0x02", "App/PC → 小车", "控制指令", "JSON动作、云台、语音和模式字段"],
        ["0x03", "小车/PC → App", "环境数据", "JSON传感器、报警和状态字段"],
    ], title="表 2-2 通信协议帧类型")
    para(doc, "报警采用统一token形式跨端传递，包括smoke、cry、close_distance、cliff、low_battery、temp_high、temp_low、light_low、light_high和light_changed等类型。App根据token展示中文报警信息，小车端根据token触发OLED状态、蜂鸣器或TTS播报，从而实现报警判定与界面展示解耦。")
    add_page_break(doc)

    heading(doc, "第3章 系统硬件设计", 1)
    para(doc, "本章按照功能组归纳硬件设计，避免逐个传感器展开过长篇幅。系统以树莓派4B为主控，通过I2C、GPIO、CSI和USB接口连接传感器与执行器，并通过传感器接口板完成接插件、电平保护和总线连接。")
    add_figure(doc, SCHEMATIC, "图 3-1 外围电路原理图", width=6.0)
    heading(doc, "3.1 树莓派主控与资源分配", 2)
    para(doc, "树莓派4B负责小车端主控和WebSocket服务，40Pin资源按功能分组规划：I2C1总线连接小车底盘控制板、OLED、PCF8591和MPU6050；GPIO连接超声波Trig/Echo、四路红外循迹和蜂鸣器；CSI连接摄像头；USB连接麦克风和声卡。")
    add_table(doc, ["资源", "连接模块", "功能说明"], [
        ["I2C1 SDA/SCL", "YB板、OLED、PCF8591、MPU6050", "电机舵机控制、显示、ADC采样和姿态采集"],
        ["GPIO16/18", "HC-SR04 Trig/Echo", "超声波测距，Echo经电阻分压接入"],
        ["GPIO27/22/17/4", "四路红外循迹", "底部悬崖检测"],
        ["GPIO12", "无源蜂鸣器", "报警提示"],
        ["CSI/USB", "摄像头、麦克风、声卡", "视频采集、语音输入和音频输出"],
    ], title="表 3-1 树莓派主控资源分配")
    heading(doc, "3.2 传感器与执行器接口", 2)
    para(doc, "测距与安全类模块包括HC-SR04超声波和四路红外循迹。HC-SR04通过Trig触发并读取Echo高电平时间计算距离，距离公式为d=(t_echo×v)/2，软件维护长度为5的滑动窗口并取中值滤波；红外循迹模块用于底部悬崖检测，当四路均为低电平时判定疑似悬空并触发强制后退。")
    para(doc, "环境感知类模块以PCF8591为核心，AIN0连接光敏电阻，AIN1连接NTC热敏电阻，AIN2连接烟雾检测输入，AIN3连接音量电位器。温度换算采用NTC B值模型：R_NTC=R_fix×ADC/(255-ADC)，T=1/(ln(R_NTC/R0)/B+1/T0)-273.15。")
    para(doc, "姿态感知类模块采用MPU6050采集三轴加速度和三轴角速度，初始化量程为±2g和±250°/s，并通过自动校准降低零偏。姿态融合采用Madgwick四元数滤波，迭代式为q_dot=1/2 q⊗omega-beta grad f，融合后的偏航角速度用于车身转向外环阻尼。")
    para(doc, "交互与显示类模块包括OLED显示屏和蜂鸣器。OLED用于显示idle、tracking、searching、sleeping和alarm等状态表情；蜂鸣器用于悬崖、距离过近等高优先级报警。音频与视觉模块包括USB麦克风、USB声卡和PiCamera2摄像头，分别承担语音/哭声输入、TTS/儿歌输出和640×480 JPEG视频采集。执行器包括双轴舵机云台和直流电机底盘，均由小车底盘控制板统一驱动。")
    heading(doc, "3.3 外围接口电路设计", 2)
    para(doc, "多个外设通过同一I2C总线连接时，需要保证地址互不冲突并控制总线线长。本系统I2C设备地址包括0x16小车底盘板、0x3C OLED、0x48 PCF8591和0x68 MPU6050。树莓派I2C1总线内置上拉，软件层面通过锁保护多线程访问，避免并发读写造成总线异常。")
    para(doc, "超声波模块Echo为5V信号，而树莓派GPIO为3.3V逻辑输入，因此采用1kΩ与2kΩ电阻分压接入GPIO18，分压后电压约3.33V，满足树莓派输入安全范围。蜂鸣器通过PWM口输出2.7kHz提示音，传感器接口板统一提供接插件与GND/VCC布线。")
    heading(doc, "3.4 硬件集成与实物展示", 2)
    para(doc, "硬件集成后，小车底盘承载树莓派、传感器接口板、摄像头云台、超声波模块、红外循迹、OLED、蜂鸣器和音频模块。实物展示章节将结合整体照片和局部照片说明模块安装位置、线束走向和运行状态。")
    add_page_break(doc)

    heading(doc, "第4章 系统软件设计", 1)
    heading(doc, "4.1 软件总体架构", 2)
    para(doc, "软件系统采用模块化和多任务协同设计。树莓派端以WebSocket服务为中心，各硬件模块继承ModuleBase线程基类，通过start、stop和stop_event统一生命周期；PC端以client.py为主循环，BackgroundService统一管理AppGateway、AsrServer和WebRtcBridge等异步服务；Android端以连接客户端和主界面Activity组织视频、环境、控制与报警显示。")
    add_figure(doc, FIG_DIR / "fig4-1-module-interaction.png", "图 4-1 软件核心模块交互时序图", width=6.0)
    add_figure(doc, FIG_DIR / "fig4-2-core-modules.png", "图 4-2 系统核心类关系图", width=6.0)
    heading(doc, "4.2 婴幼儿视觉跟随模块", 2)
    para(doc, "视觉跟随模块的数据流为：树莓派采集JPEG图像帧（640×480, quality=80）→ 0x01视频帧发送至PC → OpenCV解码为BGR → YOLO26s推理 → BabyFilter三态时序滤波 → PID双环伺服计算（内环：云台偏差到舵机增量；外环：舵机偏角到车身转向）→ 0x02指令帧回传树莓派 → 舵机与电机执行。该流程以约15~20fps的闭环速率运行，实现从图像感知到运动执行的完整视觉伺服。")
    add_figure(doc, FIG_DIR / "fig4-3-vision-follow.png", "图 4-3 视觉跟随控制流程图", width=6.0)
    heading(doc, "4.2.1 YOLO检测与BabyFilter时序滤波", 3)
    para(doc, "PC端YOLO26s模型识别baby、Adult和kids等类别，系统仅保留婴幼儿相关目标并设置置信度阈值。BabyFilter采用NONE、CANDIDATE和LOCKED三态机制：NONE状态检测到候选框后进入CANDIDATE；连续多帧匹配成功后进入LOCKED；锁定状态下通过IoU和置信度选择最佳目标，连续丢失超过阈值后退回NONE。该机制可以抑制单帧误检造成的云台和车身抖动。")
    heading(doc, "4.2.2 PID双环伺服控制", 3)
    para(doc, "内环根据目标中心与画面中心的像素偏差计算舵机角度增量，将目标尽量保持在画面中心；外环根据水平舵机偏角判断车身是否需要转向，并结合IMU偏航角速度作为阻尼项，减少车身旋转过冲。双环控制框图如图4-4所示。")
    add_figure(doc, FIG_DIR / "fig4-4-pid-loop.png", "图 4-4 PID双环控制框图", width=6.0)
    heading(doc, "4.2.3 距离跟随状态机", 3)
    para(doc, "距离跟随以超声波测距为输入。当目标在画面中心附近且无需转向时，系统根据距离区间在HOLD、FORWARD、BACKWARD和COOLDOWN状态间切换。距离过远时低速前进，距离过近时后退或禁止前进，动作结束后进入短暂冷却，避免频繁切换。")
    heading(doc, "4.3 哭声检测模块", 2)
    heading(doc, "4.3.1 YAMNet模型与滑动窗口", 3)
    para(doc, "PC端cry_detector.py基于YAMNet实现流式哭声检测。USB麦克风采集16kHz/16bit PCM音频，经WebSocket送至PC端后转换为float32波形，按约1.0s窗口执行模型推理。系统先通过RMS能量门控跳过静音段，再从YAMNet输出的521类AudioSet得分中提取婴儿哭声相关类别得分。")
    heading(doc, "4.3.2 双阈值滞后状态机", 3)
    para(doc, "哭声状态机采用高阈值进入、低阈值退出的滞后策略：哭声得分持续超过0.60并累计约2s后进入crying=true；得分低于0.40并持续约3s后退出哭声状态。低分段采用衰减累计，避免瞬间噪声或短暂停顿造成状态抖动。")
    add_figure(doc, FIG_DIR / "fig4-5-cry-state.png", "图 4-5 YAMNet哭声检测状态机图", width=6.0)
    para(doc, "与基于固定阈值或SVM分类器的哭声检测方案相比，本系统采用AudioSet大规模预训练模型结合滑动窗口和双阈值状态机，在环境噪声容忍度和误报率方面具有优势。RMS能量门控可减少静音段计算开销，低分衰减则避免瞬间噪声引起误触发。")
    heading(doc, "4.4 语音交互模块", 2)
    heading(doc, "4.4.1 系统架构", 3)
    para(doc, "语音交互模块由树莓派端麦克风采集、PC端AsrServer、百度实时语音识别服务和意图解析器组成，形成从音频采集到语义理解再到执行反馈的完整对话闭环。其数据流为：树莓派Mic → PCM(16kHz/16bit) → WebSocket(ws://PC:6006/audio) → AsrServer → 百度ASR WebSocket → 识别文本 → 意图解析 → 动作执行或TTS播报。")
    add_figure(doc, FIG_DIR / "fig4-6-voice-flow.png", "图 4-6 语音交互流程图", width=6.0)
    heading(doc, "4.4.2 百度实时ASR", 3)
    para(doc, "ASR连接建立后先发送START帧，配置dev_pid、cuid、format=pcm和sample=16000等参数；随后将麦克风PCM数据按约160ms切片持续发送；语音结束时发送FINISH帧并等待识别结果。PC端获得识别文本后不直接执行原始文本，而是交由意图解析器映射为有限动作，降低误识别造成的控制风险。")
    heading(doc, "4.4.3 意图解析与指令映射", 3)
    add_table(doc, ["指令类型", "关键词示例", "执行动作", "保持方式"], [
        ["前进", "前进、往前走、走", "电机前进", "1.5~2.2s自动归零"],
        ["后退", "后退、倒车", "电机后退", "1.5~2.2s自动归零"],
        ["左转", "左转、往左", "原地左旋", "1.5~2.2s自动归零"],
        ["右转", "右转、往右", "原地右旋", "1.5~2.2s自动归零"],
        ["停止", "停、停下来、别动", "立即停车", "单次触发"],
        ["播放儿歌", "播放儿歌、唱歌", "播放音频文件", "单次触发"],
        ["下一首", "下一首、换一首", "切换曲目", "单次触发"],
        ["停止播放", "别唱了、停下", "停止音频播放", "单次触发"],
    ], title="表 4-1 语音指令映射关系")
    para(doc, "运动类指令配置1.5s~2.2s保持时间，到期自动停止，避免因信号延迟或用户未及时补充指令导致小车持续运动引发安全风险。音频类指令采用单次触发，防止连续误触发。")
    heading(doc, "4.4.4 TTS回声抑制", 3)
    para(doc, "如果不做抑制，机器人播放的每一句语音都可能被自己的麦克风捕获，形成自己播报、自己识别、自己执行的自激回环。本文通过动态静默期策略，在播报期间丢弃ASR结果，只保留哭声检测链路。抑制时间按t_suppress=max(t_min, 2.0+len(text)×0.22)估算，使短句和长句均能覆盖主要回声时段。")
    heading(doc, "4.5 环境监测与报警", 2)
    para(doc, "环境监测模块以2Hz采样周期聚合PCF8591四通道、超声波距离、红外循迹、IMU姿态、哭声状态和系统健康状态，生成统一EnvPacket并通过0x03环境帧推送。报警策略将异常映射为alarm token，并设置TTS冷却策略：全局12s、同类45s、温度类180s。冷却只影响语音播报，不影响0x03状态推送和App提示。")
    heading(doc, "4.6 视频传输与远程通信", 2)
    para(doc, "局域网模式下，小车端以WebSocket向PC发送0x01 JPEG视频帧，PC转发至Android端显示，路径短、延迟低，适合家庭局域网使用。公网模式下，PC端通过aiortc创建WebRTC PeerConnection，经信令服务器交换offer、answer和ICE候选，结合STUN/TURN完成NAT穿透，视频轨道传输H.264画面，DataChannel同步环境数据和控制指令。")
    heading(doc, "4.7 安全保护机制", 2)
    add_table(doc, ["保护机制", "触发条件", "执行动作", "所在层级"], [
        ["距离保护", "超声波<30cm", "禁止前进，保留舵机跟踪", "小车端软件"],
        ["悬崖保护", "红外检测悬空", "立即停止并后退", "小车端软件"],
        ["通信看门狗", "0.8s未收到指令", "自动停止所有电机", "小车端软件"],
        ["麦克风断连", "2s无音频流", "自动停车", "小车端软件"],
        ["运动指令前置检查", "任何运动指令执行前", "检查安全状态，不满足则拒绝", "小车端软件"],
        ["控制权冲突", "悬崖后退中收到运动指令", "忽略指令，优先安全动作", "小车端软件"],
        ["异常报警同步", "任何异常触发", "生成alarm token并同步播报/推送", "三端协同"],
    ], title="表 4-2 安全保护机制")
    heading(doc, "4.8 Android移动端设计", 2)
    heading(doc, "4.8.1 视频显示", 3)
    para(doc, "Android端通过RaspbotConnectionClient或RaspbotWebRtcClient连接PC转发服务，接收实时视频画面并渲染显示。局域网模式接收0x01 JPEG帧解码显示，公网模式通过WebRTC视频轨道接收H.264流。")
    heading(doc, "4.8.2 环境数据面板", 3)
    para(doc, "移动端通过0x03环境数据包解析温度、光照、距离、烟雾、哭声状态和报警token等信息，在主界面以仪表盘、列表或趋势控件形式实时更新。")
    heading(doc, "4.8.3 控制面板", 3)
    para(doc, "控制面板提供前进、后退、左转、右转、停止、云台角度调整和跟随模式切换等入口。控制指令通过0x02指令帧发送至PC端，再由PC端转发至小车端执行。")
    heading(doc, "4.8.4 报警通知", 3)
    para(doc, "RaspbotAlarmService接收alarm token后在App端弹窗或系统通知中提示，报警类型包括距离过近、烟雾异常、哭声检测、悬崖风险、温度异常、光照异常和低电量等。")
    add_page_break(doc)

    heading(doc, "第5章 系统调试与功能验证", 1)
    heading(doc, "5.1 系统实物平台", 2)
    para(doc, "系统实物平台由树莓派小车、传感器接口板、双轴摄像头云台、超声波模块、红外循迹模块、OLED显示屏、蜂鸣器、USB麦克风、USB声卡和Android手机组成。整体实物和硬件模块局部效果如图5-1和图5-2所示。")
    add_figure(doc, None, "图 5-1 系统整体实物图", placeholder="待替换为图5-1：系统整体实物图")
    add_figure(doc, None, "图 5-2 硬件模块局部图", placeholder="待替换为图5-2：硬件模块局部图")
    heading(doc, "5.2 功能模块测试", 2)
    heading(doc, "5.2.1 视觉跟随测试", 3)
    para(doc, "PC端接收小车摄像头画面，对画面中的婴幼儿目标进行识别，依据目标框位置生成云台调整和距离跟随指令。当目标在画面中左右移动时，系统能够根据目标偏移调整观察方向；目标短暂离开画面或遮挡时，时序滤波策略可避免小车剧烈运动。")
    add_figure(doc, None, "图 5-3 PC端视觉识别运行截图", placeholder="待替换为图5-3：PC端视觉识别运行截图")
    para(doc, "待实测补充数据包括：目标锁定延迟约___秒，目标丢失后恢复锁定约___秒，云台跟随角度范围0°~180°，距离跟随区间26~47cm。")
    heading(doc, "5.2.2 哭声检测测试", 3)
    para(doc, "播放婴幼儿哭声样本时，YAMNet哭声得分超过0.60后约2秒内进入crying=true状态，alarm token同步至移动端和小车端；停止播放后约3秒内退出哭声状态。环境噪声如电视声、风扇声等未触发持续误报。")
    add_figure(doc, None, "图 5-4 哭声检测得分与状态切换日志", placeholder="待替换为图5-4：哭声检测得分与状态切换日志")
    heading(doc, "5.2.3 语音交互测试", 3)
    para(doc, "通过麦克风说出“前进”，系统在识别后执行前进动作并播报反馈；说出“停止”，小车立即停车。TTS播报期间再次产生的ASR结果被回声抑制机制丢弃，避免播报内容再次触发动作。")
    add_figure(doc, None, "图 5-5 语音识别文本与执行日志", placeholder="待替换为图5-5：语音识别文本与执行日志")
    heading(doc, "5.2.4 远程监控测试", 3)
    para(doc, "Android端连接PC转发服务后，能够显示实时视频画面、温度、光照、距离等环境数据，并通过控制按钮向小车发送动作指令。局域网和公网WebRTC模式均可用于远程查看，其中局域网模式更适合低延迟调试，公网模式更适合外出远程监控。")
    add_figure(doc, None, "图 5-6 Android App视频与控制界面截图", placeholder="待替换为图5-6：Android App视频与控制界面截图")
    para(doc, "待实测补充数据包括：端到端延迟约___ms（局域网）/___ms（公网WebRTC）。")
    heading(doc, "5.2.5 安全保护测试", 3)
    para(doc, "小车前方放置障碍物，距离小于30cm时前进指令被自动阻止；将小车推至桌面边缘，红外检测悬空后立即停止并后退；断开PC端连接，小车在看门狗超时时间内自动停车。上述过程说明安全保护动作优先级高于普通运动指令。")
    add_figure(doc, None, "图 5-7 报警提示与安全保护触发场景", placeholder="待替换为图5-7：报警提示与安全保护触发场景")
    heading(doc, "5.3 系统功能实现总结", 2)
    add_table(doc, ["功能模块", "展示或验证内容", "实现状态"], [
        ["实物平台搭建", "小车底盘、摄像头、传感器、显示和音频模块完成集成", "已完成"],
        ["视觉跟随", "识别目标位置并生成云台和底盘控制指令", "已完成"],
        ["哭声检测", "音频经PC端分析后同步哭声报警状态", "已完成"],
        ["语音交互", "语音指令映射为运动控制或音频播放操作", "已完成"],
        ["远程监控", "Android端显示视频、环境数据、控制按钮和报警提示", "已完成"],
        ["安全保护", "距离过近、疑似悬空、断连和报警状态下优先执行安全动作", "已完成"],
    ], title="表 5-1 系统主要功能实现情况")
    para(doc, "从功能实现情况看，系统已完成小车端、PC端和Android端之间的基本联调。视觉跟随、哭声检测、语音交互、环境监测、远程视频、安全保护等模块能够围绕统一通信协议协同工作，满足毕业设计原型系统的展示和功能验证要求。")
    add_page_break(doc)

    heading(doc, "第6章 总结与展望", 1)
    heading(doc, "6.1 工作总结", 2)
    for idx, item in enumerate([
        "完成基于移动小车的视觉主动观察系统设计，结合YOLO26s、BabyFilter和PID双环伺服控制，使系统能够围绕目标位置调整云台和车身方向。",
        "完成基于YAMNet的双阈值实时哭声检测设计，通过滑动窗口、能量门控和滞后状态机降低短时噪声造成的误触发。",
        "完成语音对话交互闭环设计与回声抑制，支持运动、停止、儿歌播放等8类语音指令，并通过动态静默期抑制TTS自激。",
        "完成局域网/公网双模远程视频通信设计，以WebSocket满足局域网低延迟查看，以WebRTC和STUN/TURN满足公网远程监控。",
        "完成多层级安全保护机制设计，覆盖障碍距离、悬崖检测、通信超时、麦克风断连、运动指令前置检查和控制权冲突管理。",
    ], 1):
        para(doc, f"（{idx}）{item}")
    heading(doc, "6.2 不足与展望", 2)
    para(doc, "本系统仍属于毕业设计原型系统，后续可从四个方面改进：第一，将YOLO和YAMNet进一步轻量化并尝试部署到边缘计算模块，降低对PC端GPU的依赖；第二，引入更稳定的多目标跟踪算法和深度传感器，提高复杂遮挡场景下的目标定位能力；第三，采集真实家庭环境下的哭声与噪声数据，训练更适合婴幼儿场景的专用检测模型；第四，完善App端历史数据、看护报告、权限认证和加密存储能力，提高长期使用的可靠性与安全性。")
    add_page_break(doc)

    heading(doc, "参考文献", 1)
    refs = [
        "[1] 韩改宁, 苏静池, 张瑞斌. 基于树莓派的智能小车的设计与开发[J]. 电子设计工程, 2024(1): 1-6.",
        "[2] 李文海, 郭伟, 宋莉. 基于树莓派4B的循迹避障小车设计[J]. 计算机与网络, 2022(19): 56-59.",
        "[3] 雒洁, 王庆坡, 周庭艳, 等. 基于树莓派和ROS系统的智能语音导盲小车[J]. 集成电路与嵌入式系统, 2023, 23(1): 50-53.",
        "[4] 庞宏鑫, 钱洪欣, 唐渝, 等. 面向婴儿安抚与监护的智能交互系统设计与实现[J]. 计算机科学与应用, 2026, 16(1): 337-352.",
        "[5] 刘任杰, 肖薇, 喻成晨, 等. 基于STM32的婴儿智能识别追踪监护系统的研究[J]. Advances in Computer and Autonomous Intelligence Research, 2024, 2(2): 1-8.",
        "[6] Liu D, Sun J M, Zheng H W. The Design and Implementation of Smart Baby Monitor System Based on ZigBee and GoAhead[J]. Applied Mechanics and Materials, 2014, 556-562: 2595-2598.",
        "[7] Raspberry Pi Foundation. Raspberry Pi 4 Model B Product Brief[EB/OL]. [2026-03-01]. https://www.raspberrypi.com/products/raspberry-pi-4-model-b/.",
        "[8] Jocher G, Chaurasia A, Qiu J. Ultralytics YOLO26[EB/OL]. [2026-03-01]. https://docs.ultralytics.com/.",
        "[9] Google Research. YAMNet: AudioSet Audio Classification Model[EB/OL]. [2026-03-01]. https://tfhub.dev/google/yamnet/1.",
        "[10] Madgwick S O H. An efficient orientation filter for inertial and inertial/magnetic sensor arrays[R]. University of Bristol, 2010: 1-24.",
        "[11] Android Developers. Android Application Development Documentation[EB/OL]. [2026-03-01]. https://developer.android.com/docs.",
        "[12] aiortc. WebRTC and ORTC implementation for Python[EB/OL]. [2026-03-01]. https://github.com/aiortc/aiortc.",
        "[13] NXP Semiconductors. PCF8591 8-bit A/D and D/A converter[EB/OL]. Product Data Sheet, 2013.",
        "[14] InvenSense. MPU-6000 and MPU-6050 Product Specification Revision 3.4[EB/OL]. 2013.",
        "[15] 百度AI开放平台. 实时语音识别技术文档[EB/OL]. [2026-03-01]. https://ai.baidu.com/tech/speech/asr.",
        "[16] Tan W, Yao Q, Liu J. Weakly Supervised Detection of Baby Cry[J]. arXiv preprint arXiv:2304.10001, 2023.",
        "[17] Fu M, Li D, Gadhiya A, et al. Infant Cry Detection Using Causal Temporal Representation[J]. arXiv preprint arXiv:2503.06247, 2025.",
        "[18] 王莹, 程倩, 郭媛媛, 等. 数智惠民——基于树莓派的智能婴儿床系统[J]. 人工智能与机器人研究, 2025, 14(1): 164-172.",
    ]
    for ref in refs:
        para(doc, ref, indent=False, size=10.5)
    add_page_break(doc)

    heading(doc, "致谢", 1)
    para(doc, "在本毕业设计完成之际，谨向所有给予我指导和帮助的人表示诚挚感谢。首先，感谢指导教师在选题确定、方案论证、系统调试和论文撰写过程中的悉心指导。老师严谨的治学态度和丰富的工程经验，使我在毕业设计过程中受益良多。")
    para(doc, "感谢学院各位老师在大学学习期间的培养和教导，使我在嵌入式系统、机器人控制、软件开发和工程实践方面得到了系统训练。感谢同学们在项目调试、资料查找和测试过程中给予的帮助。")
    para(doc, "感谢开源社区提供的树莓派、Python、OpenCV、YOLO、YAMNet、aiortc等工具和框架，为本项目实现提供了重要技术基础。最后，感谢家人的理解和支持，使我能够顺利完成毕业设计。")
    add_page_break(doc)

    heading(doc, "诚信声明", 1)
    para(doc, "本人郑重声明：所呈交的毕业设计（论文）是本人在指导教师指导下独立进行研究工作所取得的成果。除文中已经注明引用的内容外，本论文不含任何其他个人或集体已经发表或撰写过的作品成果。对本文的研究做出重要贡献的个人和集体，均已在文中以明确方式标明。本人完全意识到本声明的法律结果由本人承担。")
    para(doc, "学生签名：____________    日期：____________    指导教师签名：____________", indent=False)

    doc.save(OUTPUT_DOCX)
    print(OUTPUT_DOCX)


if __name__ == "__main__":
    build()
